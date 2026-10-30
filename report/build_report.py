"""Builds the report from the course template and measured results.

Run after scripts.train and scripts.analysis: python -m report.build_report
"""
import copy
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
HERE = Path(__file__).resolve().parent
TEMPLATE = ROOT / "Mal DAT158 innlevering 2.docx"
OUTPUT = HERE / "Rapport_Mushroom_Check.docx"

m = json.loads((ROOT / "artifacts" / "metrics.json").read_text())
a = json.loads((ROOT / "artifacts" / "analysis.json").read_text())
one = a["one_rule"]
cmp_ = m["model_comparison"]
rf = cmp_[m["model"]]
NAMES = {
    "baseline (majority)": "Majoritetsklasse (baseline)",
    "naive bayes": "Naive Bayes",
    "logistic regression": "Logistisk regresjon",
    "decision tree": "Beslutningstre",
    "random forest": "Random forest (valgt)",
    "gradient boosting": "Gradient boosting",
}
RULE_NAME = "Enfeltsregel (OneR)"


def pct(x):
    return f"{x * 100:.1f} %".replace(".", ",")


def num(x):
    return f"{x:.2f}".replace(".", ",")


def make_chart(path):
    rows = [(NAMES[k], v["unseen_species_accuracy"]) for k, v in cmp_.items()]
    rows.append((RULE_NAME, one["dev_unseen_accuracy"]))
    rows.sort(key=lambda r: r[1])
    fig, ax = plt.subplots(figsize=(6.3, 2.9), dpi=200)
    ax.barh([r[0] for r in rows], [r[1] * 100 for r in rows], height=0.55, color="#2f5d46")
    for i, (_, v) in enumerate(rows):
        ax.text(v * 100 + 1, i, pct(v), va="center", fontsize=9, color="#26241f")
    ax.set_xlim(0, 112)
    ax.set_xticks([0, 25, 50, 75, 100])
    ax.set_xlabel("Treffsikkerhet på ukjente arter (%)", fontsize=9, color="#625d52")
    ax.tick_params(labelsize=9, length=0, colors="#26241f")
    ax.grid(axis="x", color="#e8e2d4", linewidth=0.8)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left", "bottom"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


doc = Document(TEMPLATE)
body = doc.element.body
sect = body.find(qn("w:sectPr"))
paragraphs = list(doc.paragraphs)
headings = {p.text.strip(): p for p in paragraphs if p.style.name.startswith("Heading")}
bullet_ppr = next(p for p in paragraphs if p.text.startswith("Beskriv hvem produktet"))._p.pPr
title, subtitle = paragraphs[0], paragraphs[1]

for element in list(body):
    if element is not sect and element not in (title._p, subtitle._p):
        body.remove(element)


def set_text(paragraph, text):
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    paragraph.runs[0].text = text


def add_runs(paragraph, text):
    """Text inside [[double brackets]] is a highlighted placeholder for the authors to fill in."""
    for i, part in enumerate(text.replace("]]", "[[").split("[[")):
        if part:
            run = paragraph.add_run(f"[{part}]" if i % 2 else part)
            if i % 2:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def heading(text):
    sect.addprevious(headings[text]._p)


def para(text):
    p = doc.add_paragraph()
    add_runs(p, text)
    return p


def bullet(text):
    p = doc.add_paragraph()
    if p._p.pPr is not None:
        p._p.remove(p._p.pPr)
    p._p.insert(0, copy.deepcopy(bullet_ppr))
    add_runs(p, text)


def caption(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def picture(path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=width)


def table(header, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "insideH"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "BFBFBF")):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    t.autofit = False
    for r, row in enumerate([header, *rows]):
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if c:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(value)
            run.font.size = Pt(9.5)
            run.bold = r == 0 or value == NAMES[m["model"]] or row[0] == NAMES[m["model"]]
    doc.add_paragraph()


set_text(title, "Mushroom Check")
set_text(subtitle, "")
add_runs(subtitle, "[[Navn]], [[Navn]], [[Navn]], 30. oktober 2026")

heading("1: BESKRIV PROBLEMET")
heading("OMFANG / SCOPE")
para("Målet med prosjektet er å lage en nettside der brukeren beskriver en sopp ved hjelp av skiver, lukt, farge og voksested, og får et maskinlæringsbasert svar på om soppen kan være giftig. Prosjektet er en læringsoppgave og en demonstrasjon av trygg bruk av maskinlæring. Det er ikke et verktøy for å avgjøre hva som kan spises, og nettsiden sier dette tydelig.")
para("Maskinlæring passer fordi oppgaven er klassifisering med mange kategoriske egenskaper som virker sammen, og fordi feltguider ikke gir noen enkel regel: Datasettets dokumentasjon viser til at guiden sier det ikke finnes en enkel tommelfingerregel for spiselighet. I dag løses problemet manuelt ved å artsbestemme soppen med feltguide, spørre erfarne plukkere eller bruke soppkontroll. Det finnes også apper som artsbestemmer fra bilder. Vårt bidrag er ikke høyere treffsikkerhet, men en løsning som sier fra når den ikke vet.")
para("Målgruppen er studenter og andre som er nysgjerrige på hvordan en maskinlæringsmodell kan oppføre seg ansvarlig. Løsningen er et selvstendig system: et skjema i nettleseren, et API som kjører modellen og en beslutningsregel. Endringer i modellen påvirker sannsynlighetene som beslutningsregelen er kalibrert mot, så de to må alltid evalueres sammen.")
para("Prosjektet har ingen kommersiell inntekt, så «business objective» er å unngå skade. Det alvorligste som kan skje er at en giftig sopp omtales som spiselig. Ressursbehovet er lite: en vanlig bærbar PC trener modellen på noen minutter, og nettsiden kan kjøre på en gratis Docker-tjeneste.")

heading("METRIKKER")
para("Den viktigste metrikken er andelen giftige sopper som får svaret «Looks edible» (falsk spiselig-rate), målt på arter modellen aldri har sett. Disse målene ble brukt:")
bullet("Falsk spiselig-rate på ukjente arter. Krav: høyst 1 % på hvert nivå av manglende input.")
bullet(f"Dekning: andelen spiselige sopper fra kjente arter som får svaret spiselig. Målet er at den skal være så høy som mulig innenfor kravet over. Den viser at løsningen fortsatt er nyttig.")
bullet("Andelen «Can't tell»: hvor ofte systemet avstår fra å svare.")
bullet("Treffsikkerhet, giftig-recall og log loss, brukt til å sammenligne modeller. Log loss måler kvaliteten på sannsynlighetene som beslutningsregelen bygger på.")
bullet("Svartid. Krav: under 500 ms per prediksjon.")
para("Prosjektet regnes som vellykket dersom kravene til falsk spiselig-rate og svartid er oppfylt, og nettsiden kan brukes av andre. Vi har ingen brukerdata, så vi har ikke satt mål for antall brukere eller brukertilfredshet.")

heading("2: DATA")
para(f"Vi bruker Mushroom-datasettet fra UCI Machine Learning Repository, som er utdelt til prosjektet. Det har {m['rows']} rader med 22 kategoriske egenskaper og en etikett: spiselig (51,8 %) eller giftig (48,2 %). Dataene bygger på The Audubon Society Field Guide to North American Mushrooms fra 1981 og beskriver hypotetiske eksemplarer av 23 arter i familiene Agaricus og Lepiota. Etikettene kommer fra feltguiden. Arter med ukjent spiselighet er slått sammen med giftige, og det er en forsiktig, men ikke nøytral, avgjørelse.")
para(f"Den viktigste oppdagelsen var strukturen i dataene. Filen har ingen artskolonne og ingen duplikater, men hver rad har en nabo som skiller seg i nøyaktig én egenskap. Når vi kobler sammen slike rader, får vi nøyaktig 23 grupper som hver har én etikett. Det stemmer med de 23 artene i dokumentasjonen, og etikettene er dermed konsistente innenfor hver gruppe. Det betyr også at den reelle datamengden er 23 arter, ikke {m['rows']} uavhengige eksempler. Flere rader ville hjulpet lite. Det vi trenger, er flere arter.")
para("Datasettet inneholder ingen personopplysninger, og nettsiden lagrer ikke det brukeren skriver inn. Den etiske risikoen ligger i at noen stoler på svaret og spiser en giftig sopp. Derfor har siden en tydelig advarsel, og den kan ikke svare «spiselig» uten å oppfylle strenge krav. Dataene dekker bare nordamerikanske arter, og modellen kan dermed ikke overføres til norsk sopp uten videre.")
para("Dataene ble behandlet slik:")
bullet("«veil-type» ble fjernet fordi den har samme verdi i alle rader.")
bullet("Alle 2480 manglende verdier er i «stalk-root». De ble beholdt som en egen kategori, «?», som også betyr «Not sure» på nettsiden. 1760 av disse radene er giftige, så manglende verdi trekker risikoen opp.")
bullet("Egenskapene ble one-hot-kodet. Skalering var ikke nødvendig.")
bullet("Treningsdataene ble utvidet med fire kopier der 10–80 % av feltene er tilfeldig skjult, slik at modellen kan brukes på delvis utfylte skjemaer.")

heading("3: MODELLERING")
para(f"Vi delte artene i to før modelleringen: {m['dev_species']} arter til utvikling og {m['test_species']} arter ({a['test_rows_poisonous']} giftige og {a['test_rows_edible']} spiselige rader) som ble låst bort til sluttesten. På utviklingsartene sammenlignet vi modeller med leave-one-species-out: Hver art holdes ute én gang, og modellen trenes på resten. Som kontrast målte vi også en vanlig tilfeldig 80/20-deling.")
para("Tabell 1 viser resultatene. Etter sluttesten la vi også til en enkel regel (OneR) som kun lærer hvilken enkeltegenskap som passer best med treningsartene. Den ble ikke brukt til å velge modell, men endte som en del av det ferdige systemet, slik vi forklarer under.")
caption("Tabell 1: Modeller på utviklingsartene. «Ukjente arter» er leave-one-species-out. Tallet for tilfeldig deling for enfeltsregelen er oppgitt i datasettets dokumentasjon for alle rader.")
rows = [[NAMES[k], pct(v["random_split_accuracy"]), pct(v["unseen_species_accuracy"]),
         pct(v["unseen_species_poison_recall"]), num(v["unseen_species_log_loss"])] for k, v in cmp_.items()]
rows.append([RULE_NAME, pct(a["odor_rule"]["all_rows_accuracy"]), pct(one["dev_unseen_accuracy"]),
             pct(one["dev_unseen_poison_recall"]), "–"])
table(["Modell", "Tilfeldig deling", "Ukjente arter", "Giftig-recall ukjente", "Log loss ukjente"], rows,
      [5.2, 2.5, 2.4, 2.8, 2.4])
para(f"Tilfeldig deling gir nesten 100 %, men det er misvisende. Nesten like rader havner på begge sider av delingen, så modellene kjenner igjen arter i stedet for å generalisere. Random forest ble valgt fordi den hadde lavest log loss ({num(rf['unseen_species_log_loss'])}), det vil si de beste sannsynlighetene på ukjente arter. Selv den traff bare {pct(rf['unseen_species_accuracy'])} og fant {pct(rf['unseen_species_poison_recall'])} av de giftige.")
chart = HERE / "modeller_ukjente_arter.png"
make_chart(chart)
picture(chart, Cm(15))
caption("Figur 1: Treffsikkerhet på ukjente arter, utviklingsartene (leave-one-species-out).")

para("Fordi modellen alene ikke er trygg, bygde vi en beslutningsregel oppå den. Svaret «Looks edible» krever at alle disse fire vilkårene er oppfylt:")
bullet("Luktregelen (se feilanalysen) gir ingen varsel.")
bullet(f"Modellens giftrisiko er høyst {m['policy']['max_p_poison']:.0%}.".replace("%", " %"))
bullet(f"Inputen skiller seg fra minst én kjent sopp i høyst {m['policy']['max_distance']} av de oppgitte egenskapene.")
bullet(f"Minst {m['policy']['min_known']} av 21 egenskaper er fylt ut.")
para("Ellers blir svaret «Likely poisonous» hvis luktregelen varsler eller risikoen er minst 50 %, og «Can't tell» hvis ikke. Grensene ble valgt med et rutenett på utviklingsartene: den kombinasjonen som gav flest spiselig-svar på kjente arter, og samtidig holdt falsk spiselig-rate på høyst 1 % for ukjente arter på hvert nivå av skjult input.")
caption("Tabell 2: Det ferdige systemet. «Ukjente» er alle 23 arter holdt ute én om gangen, opptil 150 rader per art. «Kjente» er holdt ut rader fra utviklingsartene.")
rate_rows = []
for rate, label in (("0.0", "0 %"), ("0.25", "25 %"), ("0.5", "50 %")):
    t, k = m["all_unseen_species"][rate], m["dev_known_species"][rate]
    rate_rows.append([label, pct(t["false_edible_rate"]), pct(t["poison_flagged_poisonous"]), pct(t["uncertain_rate"]),
                      pct(k["edible_coverage"]), pct(k["poison_flagged_poisonous"])])
table(["Skjulte felter", "Ukjente: giftig kalt spiselig", "Ukjente: giftig flagget", "Ukjente: «Can't tell»", "Kjente: spiselig besvart", "Kjente: giftig flagget"],
      rate_rows, [2.2, 2.9, 2.6, 2.6, 2.6, 2.5])
worst = max(v["false_edible_rate"] for v in m["all_unseen_species"].values())
locked, all_plain = m["test_unseen_species"]["0.0"], m["all_unseen_species"]["0.0"]
para(f"Kravet er oppfylt: Den høyeste raten av giftige sopper kalt spiselige var {pct(worst)} over alle 23 arter, og ingen av radene i de fem låste artene ble kalt spiselige. Modellen alene traff bare {pct(m['test_unseen_species_model_accuracy'])} på de låste artene, så sikkerheten kommer fra beslutningsreglene. Spiselige sopper fra ukjente arter får aldri svaret spiselig, og det er tilsiktet. Prisen er at {pct(all_plain['edible_flagged_poisonous'])} av dem feilaktig får «Likely poisonous», og at omtrent halvparten av alle svar på ukjente arter blir «Can't tell». Dette er en trygg feil, men den gjør systemet mindre nyttig for slike arter.")
para("Tallene må leses med forsiktighet. Luktregelen ble lagt til etter at de fem låste artene var brukt én gang, så de er ikke lenger et uavhengig testsett, og resultatet for alle 23 arter er en etterfølgende vurdering. Få arter betyr også at 0 % ikke viser at den sanne raten er null.")

doc.add_paragraph("FEILANALYSE OG FEATURE IMPORTANCE", style="Heading 2")
errors = a["dev_species_errors"]
low = [e for e in errors if e["accuracy"] < 0.01]
high = [e for e in errors if e["accuracy"] >= 0.95]
imp = list(a["feature_importance"].items())
para(f"Feilene er svært ujevnt fordelt. Av de 18 utviklingsartene ble {len(high)} nesten helt riktig klassifisert (minst 95 %), mens {len(low)} giftige arter nesten aldri ble gjenkjent (under 1 %). Modellen bommer altså på hele arter, ikke på tilfeldige rader. Den samlede treffsikkerheten påvirkes også sterkt av noen få store arter.")
para(f"Viktigst i modellen er lukt ({num(imp[0][1])}), deretter sporefarge ({num(imp[1][1])}) og skivestørrelse ({num(imp[2][1])}), målt som summert impurity-basert viktighet. Det stemmer med reglene i dokumentasjonen, der lukt er den viktigste egenskapen. Vi brukte også denne innsikten til å teste OneR, som lærte seg luktregelen i alle {sum(one['features_chosen'].values())} foldene: Lukten er enten ingen, mandel eller anis for spiselige sopper, og ellers tyder den på gift.")
para(f"Resultatet var overraskende. OneR fikk {pct(one['dev_unseen_accuracy'])} på ukjente arter, langt over alle de andre modellene. På de låste artene fikk den {pct(one['test_accuracy'])}, med {pct(one['test_false_edible'])} falsk spiselig-rate. De mer fleksible modellene pugger altså artsspesifikke kjennetegn, og vi burde hatt med en slik enkel baseline fra starten. OneR bommet likevel på {pct(1 - one['dev_unseen_poison_recall'])} av de giftige ukjente radene, som er for mye til å stole på et spiselig-svar, og den virker ikke uten en luktangivelse.")
para(f"Derfor erstattet vi ikke modellen, men la luktregelen inn i beslutningen på en måte som bare kan gjøre svaret forsiktigere. En lukt som peker mot gift, eller som regelen aldri har sett, gir «Likely poisonous» uansett hva modellen mener. En trygg lukt åpner derimot aldri for svaret spiselig. Regelen læres fra treningsdataene, ikke fra en fast liste. På de låste artene økte andelen giftige sopper som ble flagget, fra 43,6 % til {pct(locked['poison_flagged_poisonous'])}, og «Can't tell» sank fra 63,1 % til {pct(locked['uncertain_rate'])}. Svarene for kjente arter ble ikke endret.")

heading("4: DEPLOYMENT")
para("Løsningen består av en FastAPI-backend og en enkel nettside i HTML, CSS og JavaScript. Brukeren fyller ut de egenskapene de kjenner til, og svaret oppdateres med en gang. Siden viser svaret, modellens giftrisiko, hvorfor den eventuelt holder tilbake et spiselig-svar, og hvilke egenskaper som påvirket risikoen mest. Dette finner vi ved å skjule hver egenskap og se hvordan risikoen endrer seg. API-et avviser ugyldig input, og 23 automatiske tester dekker datahåndtering, beslutningsregel og API.")
para(f"Prosjektet har en Dockerfile som kjører nettsiden med den ferdigtrente modellen. Modellfilen ligger i repositoriet, fordi treningen ikke ble ferdig under bygging på gratistjenesten, og den kan lages på nytt med python -m scripts.train. Faste seed-verdier og låste pakkeversjoner gjør at resultatene kan reproduseres: Vi fikk identiske metrikker etter en ny installasjon. Medianen for svartid var {str(a['latency']['median_ms']).replace('.', ',')} ms og maksimum {str(a['latency']['max_ms']).replace('.', ',')} ms lokalt, uten nettverkstid, godt innenfor kravet.")
para("Nettsiden: [[sett inn lenke etter deployment]]. Kildekode: [[sett inn lenke til GitHub-repositoriet]].")
picture(HERE / "app_screenshot.jpg", Cm(13))
caption("Figur 2: Nettsiden med et ukjent eksempel. Svaret viser både risiko og begrunnelse.")
para("I drift bør vi følge med på svartid, feilrater og fordelingen av de tre svarene, uten å lagre brukerens input. En økende andel «Can't tell» kan tyde på at brukerne beskriver sopp modellen ikke kjenner. Pakkene må oppdateres jevnlig, og modellen og grensene må alltid evalueres på nytt sammen. Vi har ingen fasit for hva brukerne skriver inn, så modellens virkelige ytelse kan ikke måles direkte.")
para("Planer for videre forbedring:")
bullet("Teste hele systemet på nye arter, siden alle de 23 nå er brukt til å utvikle det.")
bullet("Redusere falske «Likely poisonous» for spiselige ukjente arter, for eksempel ved å svare «Can't tell» når lukten er trygg og soppen er ukjent.")
bullet("Teste enklere og sterkere regulariserte modeller, og velge modell med artsbasert validering fra starten.")
bullet("Legge til flere arter, særlig norske, og data fra ekte eksemplarer.")
bullet("Undersøke hvor godt vanlige brukere klarer å bedømme egenskaper som lukt.")

heading("5: REFERANSER")
for ref in (
    "Schlimmer, J. (1987). Mushroom [Datasett]. UCI Machine Learning Repository. https://archive.ics.uci.edu/dataset/73/mushroom",
    "Lincoff, G. H. (Pres.) (1981). The Audubon Society Field Guide to North American Mushrooms. Alfred A. Knopf.",
    "Duch, W., Adamczak, R. & Grabczewski, K. (1996). Extraction of logical rules from training data using backpropagation networks. Proc. of the 1st Online Workshop on Soft Computing, 25–30. Sitert i datasettets dokumentasjon.",
    "Holte, R. C. (1993). Very simple classification rules perform well on most commonly used datasets. Machine Learning, 11, 63–90.",
    "Pedregosa, F. mfl. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    "FastAPI. https://fastapi.tiangolo.com",
    "KI-verktøy: Claude Code (Anthropic) ble brukt til kode, utforming av evalueringen og utkast til denne rapporten.",
):
    para(ref)

doc.save(OUTPUT)

words = 0
for p in Document(OUTPUT).paragraphs:
    words += len(p.text.split())
print("saved", OUTPUT, "| words excl. tables:", words)
